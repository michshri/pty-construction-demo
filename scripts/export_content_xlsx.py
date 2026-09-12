import json
import os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

root = r"c:\Users\Michael\Desktop\Cursor\pty-construction-demo\pty-construction-demo"
output_file = os.path.join(root, "Vanguardia_reusable_content.xlsx")

files = [
    ("Home Page", os.path.join(root, "content", "pages", "home.json")),
    ("Services", os.path.join(root, "content", "tiles", "services.json")),
    ("Segments", os.path.join(root, "content", "tiles", "segments.json")),
    ("Why Vanguardia", os.path.join(root, "content", "reusable", "why-vanguardia.json")),
    ("Process", os.path.join(root, "content", "reusable", "process.json")),
    ("Projects", os.path.join(root, "content", "projects", "projects.json")),
    ("Site Settings", os.path.join(root, "content", "settings", "site-settings.json")),
    ("EN Locale", os.path.join(root, "locales", "en", "common.json")),
]


def flatten_object(prefix, value, rows):
    if isinstance(value, dict):
        if all(isinstance(v, (str, int, float, bool)) or v is None for v in value.values()):
            rows.append({"key": prefix, "en": value.get("en"), "es": value.get("es"), "fr": value.get("fr"), "raw": json.dumps(value, ensure_ascii=False)})
            return
        for key, child in value.items():
            child_prefix = f"{prefix}.{key}" if prefix else key
            flatten_object(child_prefix, child, rows)
    elif isinstance(value, list):
        for idx, item in enumerate(value):
            flatten_object(f"{prefix}[{idx}]", item, rows)
    else:
        rows.append({"key": prefix, "en": value, "es": "", "fr": "", "raw": ""})


def add_sheet(wb, title, data):
    ws = wb.create_sheet(title=title[:31])
    ws.freeze_panes = "A2"
    headers = ["Key", "Type", "EN", "ES", "FR", "Notes"]
    ws.append(headers)
    header_fill = PatternFill("solid", fgColor="D9EAF7")
    for cell in ws[1]:
        cell.font = Font(bold=True, color="1F1F1F")
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for item in data:
        row = [
            item.get("key", ""),
            item.get("type", ""),
            item.get("en", ""),
            item.get("es", ""),
            item.get("fr", ""),
            item.get("notes", ""),
        ]
        ws.append(row)

    for column_cells in ws.columns:
        max_length = 0
        column = column_cells[0].column_letter
        for cell in column_cells:
            if cell.value is not None:
                max_length = max(max_length, len(str(cell.value)))
        ws.column_dimensions[column].width = min(max(12, max_length + 2), 60)

    for row in ws.iter_rows():
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)


def normalize_data_from_json(json_path):
    with open(json_path, "r", encoding="utf-8") as file:
        data = json.load(file)

    rows = []
    flattened = []
    flatten_object("root", data, flattened)

    for entry in flattened:
        key = entry["key"].replace("root.", "")
        if key.startswith("root"):
            key = key.replace("root", "", 1).lstrip(".")
        if key.startswith("["):
            key = key[1:]

        if key and not key.startswith("root"):
            key = key.replace("[0]", "")
            key = key.replace(".", " / ")
        else:
            key = key or "root"

        value = json.loads(entry.get("raw", "{}")) if entry.get("raw") else {}
        if isinstance(value, dict):
            en = value.get("en", "")
            es = value.get("es", "")
            fr = value.get("fr", "")
            notes = ""
            row_type = "locale_object"
        else:
            en = entry.get("en", "")
            es = entry.get("es", "")
            fr = entry.get("fr", "")
            notes = ""
            row_type = "value"

        rows.append({
            "key": key,
            "type": row_type,
            "en": en,
            "es": es,
            "fr": fr,
            "notes": notes,
        })

    return rows


wb = Workbook()
ws = wb.active
ws.title = "Instructions"
ws.append(["Vanguardia Reusable Content Review"])
ws.append([])
ws.append(["Instructions:"])
ws.append(["1. Review and edit the values in each sheet."])
ws.append(["2. Keep the structure for EN / ES / FR translations as needed."])
ws.append(["3. After editing, use this workbook as the source to update the JSON content."])
ws.append([])
ws.append(["Generated from:"])
for _, path in files:
    ws.append([path])

for cell in ws[1]:
    cell.font = Font(bold=True, size=14)
for row in ws.iter_rows(min_row=3):
    for cell in row:
        cell.alignment = Alignment(vertical="top", wrap_text=True)

for sheet_name, path in files:
    try:
        data = normalize_data_from_json(path)
    except Exception as exc:
        data = [{"key": "ERROR", "type": "error", "en": str(exc), "es": "", "fr": "", "notes": ""}]
    add_sheet(wb, sheet_name, data)

# Move Instructions to front
instructions = wb["Instructions"]
wb._sheets.remove(instructions)
wb._sheets.insert(0, instructions)

wb.save(output_file)
print(f"Created workbook: {output_file}")
